#!/usr/bin/env python3
"""Publish the decks dropped in slideshow/inbox/ to the /slideshow/ tool.

Each deck is a PowerPoint file (.pptx) and/or a PDF exported from PowerPoint (.pdf) sharing the
same file name, with an optional <name>.json metadata file:
  {"id": "...", "title": "...", "description": "...", "language": "en", "date": "YYYY-MM-DD", "tags": ["..."]}

Slides are rendered from the PDF when one is provided (pixel-perfect), otherwise the .pptx is
converted with LibreOffice. Titles, chapters and visible slide text come from the .pptx when
present (speaker notes are never read), otherwise from the PDF text. Every visible slide is
written as WebP (full size + thumbnail) to decks/<id>/ and the deck is added to or replaced in
decks.json. Published inbox files are deleted unless --keep is passed.

Requires: Python 3 with Pillow and python-pptx, poppler-utils (pdftoppm, pdftotext) and
LibreOffice (only for .pptx files without a matching PDF).
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

from PIL import Image

ROOT = Path(__file__).resolve().parent
INBOX = ROOT / "inbox"
DECKS = ROOT / "decks"
MANIFEST = ROOT / "decks.json"
WIDTH, THUMB_WIDTH = 1600, 360
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ID_PATTERN = re.compile(r"^[a-z0-9][a-z0-9-]*$")
DATE_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class DeckError(Exception):
    pass


def slugify(value):
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-") or "deck"


def clean(value):
    return re.sub(r"\s+", " ", value or "").strip()


def clip(value, limit):
    return value if len(value) <= limit else value[: limit - 3].rstrip() + "..."


def run(*args):
    result = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.returncode:
        raise DeckError(f"{Path(args[0]).name} failed: {(result.stderr or result.stdout).strip()[-800:]}")
    return result.stdout


def find_soffice():
    for name in ("soffice", "libreoffice"):
        if shutil.which(name):
            return shutil.which(name)
    for path in (r"C:\Program Files\LibreOffice\program\soffice.exe", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if Path(path).exists():
            return path
    raise DeckError("LibreOffice is required to convert .pptx files (or add a PDF exported from PowerPoint)")


def group_inbox():
    groups = {}
    for path in sorted(INBOX.iterdir()) if INBOX.exists() else []:
        ext = path.suffix.lower()
        if path.is_file() and ext in (".pptx", ".pdf", ".json"):
            groups.setdefault(path.stem, {})[ext[1:]] = path
    return {stem: files for stem, files in groups.items() if "pptx" in files or "pdf" in files}


def ensure_not_encrypted(path):
    with path.open("rb") as handle:
        if handle.read(8) == OLE_MAGIC:
            raise DeckError(
                f"{path.name} is encrypted (password or sensitivity label with protection). "
                "Remove the password and apply a sensitivity label without encryption in PowerPoint, then drop it again."
            )


def ensure_pdf_not_protected(path, name):
    if "Microsoft Information Protection" in clean(run("pdftotext", "-l", "1", "-enc", "UTF-8", str(path), "-")):
        raise DeckError(
            f"{name} is protected by Microsoft Information Protection (sensitivity label with encryption). "
            "Apply a non-encrypted sensitivity label in PowerPoint, then export or drop the deck again."
        )


def pptx_texts(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(shape.shapes)
            elif shape.has_text_frame and shape.text_frame.text.strip():
                yield clean(shape.text_frame.text)

    deck = Presentation(str(path))
    slides = [list(walk(slide.shapes)) for slide in deck.slides if slide._element.get("show") not in ("0", "false")]
    return slides, deck.core_properties


def pdf_texts(pdf, count):
    slides = []
    for page in range(1, count + 1):
        text = run("pdftotext", "-f", str(page), "-l", str(page), "-enc", "UTF-8", str(pdf), "-")
        slides.append([clean(line) for line in text.splitlines() if clean(line)])
    return slides


def describe(slides_texts):
    slides, chapters = [], []
    for number, texts in enumerate(slides_texts, start=1):
        first = texts[0] if texts else ""
        title = first
        if re.fullmatch(r"\d{1,2}|[A-Z]", first) and len(texts) >= 3:
            title = texts[2]
            chapters.append({"label": first, "title": texts[2], "slide": number})
        elif first.startswith("//") and len(texts) >= 2:
            title = texts[1]
        slides.append({"title": clip(title, 90), "text": clip(" · ".join(texts), 600)})
    return slides, chapters


def metadata(stem, files, core):
    meta = {}
    if "json" in files:
        try:
            meta = json.loads(files["json"].read_text(encoding="utf-8-sig"))
        except json.JSONDecodeError as error:
            raise DeckError(f"{files['json'].name} is not valid JSON: {error}") from error
    tags = meta.get("tags")
    if tags is None:
        tags = re.split(r"[,;]", core.keywords) if core and core.keywords else []
    elif isinstance(tags, str):
        tags = tags.split(",")
    result = {
        "id": meta.get("id") or slugify(stem),
        "title": clean(meta.get("title") or (core and core.title) or stem.replace("-", " ").replace("_", " ")),
        "description": clean(meta.get("description") or (core and (core.subject or core.comments)) or ""),
        "language": meta.get("language") or "en",
        "author": clean(meta.get("author") or "Benoit Gaumard"),
        "date": meta.get("date") or dt.date.today().isoformat(),
        "tags": [clean(tag) for tag in tags if clean(tag)],
    }
    if not ID_PATTERN.fullmatch(result["id"]):
        raise DeckError(f"Invalid deck id {result['id']!r}: use lowercase letters, digits and dashes")
    if not DATE_PATTERN.fullmatch(result["date"]):
        raise DeckError(f"Invalid date {result['date']!r}: use YYYY-MM-DD")
    return result


def publish(stem, files, work):
    if "pptx" in files:
        ensure_not_encrypted(files["pptx"])
    pdf = files.get("pdf")
    if not pdf:
        source = work / "deck.pptx"
        shutil.copyfile(files["pptx"], source)
        run(find_soffice(), "--headless", "--norestore", f"-env:UserInstallation=file:///{(work / 'profile').as_posix().lstrip('/')}",
            "--convert-to", "pdf", "--outdir", str(work), str(source))
        pdf = work / "deck.pdf"
        if not pdf.exists():
            raise DeckError(f"LibreOffice could not convert {files['pptx'].name}")
    ensure_pdf_not_protected(pdf, files.get("pdf", files.get("pptx")).name)

    pages = work / "pages"
    pages.mkdir()
    run("pdftoppm", "-png", "-scale-to-x", str(WIDTH * 2), "-scale-to-y", "-1", str(pdf), str(pages / "p"))
    images = sorted(pages.glob("p-*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[1]))
    if not images:
        raise DeckError(f"No slide rendered for {stem}")

    texts, core = None, None
    if "pptx" in files:
        texts, core = pptx_texts(files["pptx"])
        if len(texts) != len(images):
            print(f"  ! {stem}: {len(texts)} visible slides in the .pptx but {len(images)} pages rendered, using PDF text")
            texts = None
    if texts is None:
        texts = pdf_texts(pdf, len(images))
    slides, chapters = describe(texts)
    meta = metadata(stem, files, core)

    target = DECKS / meta["id"]
    if target.exists():
        shutil.rmtree(target)
    target.mkdir(parents=True)
    height = WIDTH * 9 // 16
    for number, png in enumerate(images, start=1):
        with Image.open(png) as image:
            image = image.convert("RGB")
            height = round(image.height * WIDTH / image.width)
            image.resize((WIDTH, height), Image.LANCZOS).save(target / f"{number:03d}.webp", "WEBP", quality=84, method=6)
            image.resize((THUMB_WIDTH, round(image.height * THUMB_WIDTH / image.width)), Image.LANCZOS).save(
                target / f"{number:03d}-thumb.webp", "WEBP", quality=72, method=6)
    return {**meta, "width": WIDTH, "height": height, "slideCount": len(images), "chapters": chapters, "slides": slides}


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--keep", action="store_true", help="keep the inbox files after publishing")
    args = parser.parse_args()

    groups = group_inbox()
    if not groups:
        print("No deck found in slideshow/inbox/.")
        return 0
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8")) if MANIFEST.exists() else {"decks": []}
    published, failed = [], []
    for stem, files in groups.items():
        print(f"Publishing {stem} ({', '.join(sorted(files))})")
        try:
            with tempfile.TemporaryDirectory(prefix="slideshow-") as work:
                entry = publish(stem, files, Path(work))
        except Exception as error:  # keep publishing the other decks
            failed.append(f"{stem}: {error}")
            print(f"  x {error}", file=sys.stderr)
            continue
        manifest["decks"] = [deck for deck in manifest["decks"] if deck["id"] != entry["id"]] + [entry]
        published.append(entry)
        print(f"  ok {entry['id']}: {entry['slideCount']} slides, {len(entry['chapters'])} chapters")
        if not args.keep:
            for path in files.values():
                path.unlink()

    if published:
        manifest["decks"].sort(key=lambda deck: (deck["date"], deck["id"]), reverse=True)
        manifest = {"generatedAt": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "decks": manifest["decks"]}
        MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")

    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary:
        with open(summary, "a", encoding="utf-8") as handle:
            handle.write("### Slideshow decks\n\n")
            for entry in published:
                handle.write(f"- ✅ [{entry['title']}](https://benoit-gaumard.io/slideshow/?deck={entry['id']}): {entry['slideCount']} slides\n")
            for message in failed:
                handle.write(f"- ❌ {message}\n")
    output = os.environ.get("GITHUB_OUTPUT")
    if output:
        with open(output, "a", encoding="utf-8") as handle:
            handle.write(f"published={' '.join(entry['id'] for entry in published)}\n")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
